"""
HCRE Wellness Center – Cash Flow Dashboard & Actuals Tracker
Combination of interactive dashboard + actuals entry.
"""

import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots
from pathlib import Path
from copy import deepcopy
from datetime import datetime
import io

# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------
st.set_page_config(
    page_title="HCRE Cash Flow",
    page_icon="🏗️",
    layout="wide",
    initial_sidebar_state="expanded",
)

# Excel path – works for local runs and Streamlit Community Cloud (flat repo root)
_APP_DIR = Path(__file__).resolve().parent
_CANDIDATES = [
    _APP_DIR / "Cash Flow Forecast.xlsx",                       # preferred: same folder as app.py (Cloud / flat layout)
    _APP_DIR.parent / "Cash Flow Forecast.xlsx",                # unzip layout: Excel next to hcre_cashflow_app/
    Path("/home/workdir/attachments/Cash Flow Forecast.xlsx"),  # original sandbox path
]
SOURCE_XLSX = next((p for p in _CANDIDATES if p.exists()), None)
if SOURCE_XLSX is None:
    raise FileNotFoundError(
        "Could not find 'Cash Flow Forecast.xlsx'. "
        "Place it in the same folder as app.py (recommended for Streamlit Cloud)."
    )
FACILITY_LIMIT = 10_000_000  # RMH Mezzanine capital facility

MONTH_LABELS = [
    "Jul-26", "Aug-26", "Sep-26", "Oct-26",
    "Nov-26", "Dec-26", "Jan-27", "Feb-27",
    "Mar-27", "Apr-27", "May-27", "Jun-27", "Jul-27",
]
STATUS = ["Actual"] * 4 + ["Forecast"] * 9  # Jul–Oct actual, rest forecast

# Line items that are editable (leaf nodes, not category totals)
EDITABLE_ROWS = {
    13: "Designs & 3D renders (ARC wellness 1)",
    14: "Designs & 3D renders (ARC wellness 2)",
    15: "Designs – layouts, site plan, elevations (ARC BKA 1)",
    16: "Municipal drawing preparation & Approval (LVG 10)",
    22: "Johan Erwee (Engineer 3147)",
    25: "Land Surveyor",
    26: "Attorney",
    27: "Auditors",
    28: "PM",
    31: "NHBRC",
    32: "Plan Approval",
    33: "Sewage",
    34: "Water",
    35: "Electrical",
    36: "Storm Water",
    37: "Rates and Taxes",
    40: "Site Clearance",
    41: "Fencing",
    42: "Security",
    43: "Casual Labour",
}

CATEGORY_PARENTS = {
    "Professional fees – Architect": [13, 14, 15, 16],
    "Professional fees – Engineer": [22],
    "Professional fees – Other": [25, 26, 27, 28],
    "Municipal fees & contributions": [31, 32, 33, 34, 35, 36, 37],
    "Site preparation & earthworks": [40, 41, 42, 43],
}


# ---------------------------------------------------------------------------
# Data loading & helpers
# ---------------------------------------------------------------------------
@st.cache_data
def load_raw_excel(path: str) -> dict:
    """Load key sheets/sections from the source workbook into plain dicts."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active

    # Monthly cost values for every relevant row
    costs = {}
    budgets = {}
    for row in range(12, 62):
        name = ws.cell(row=row, column=1).value
        if not name:
            continue
        name = str(name).strip()
        budget = ws.cell(row=row, column=2).value
        monthly = []
        for col in range(7, 20):
            v = ws.cell(row=row, column=col).value
            monthly.append(float(v) if v is not None else 0.0)
        costs[row] = {"name": name, "monthly": monthly}
        budgets[row] = float(budget) if budget is not None else None

    # Interest rates per month
    rates = []
    for col in range(7, 20):
        r = ws.cell(row=75, column=col).value
        rates.append(float(r) if r is not None else 0.1175)

    # Revenue budget
    life_right_budget = ws.cell(row=7, column=2).value or 0

    return {
        "costs": costs,
        "budgets": budgets,
        "rates": rates,
        "life_right_budget": float(life_right_budget),
    }


def init_session_state():
    if "data" not in st.session_state:
        raw = load_raw_excel(str(SOURCE_XLSX))
        st.session_state.data = raw
        st.session_state.entries = []  # list of {row, month_idx, amount, note, ts}
        st.session_state.overrides = {}  # (row, month_idx) -> amount  (overrides original)


def get_monthly_costs() -> dict:
    """Return {row: list of 13 monthly amounts} applying any overrides."""
    costs = deepcopy(st.session_state.data["costs"])
    for (row, midx), amt in st.session_state.overrides.items():
        if row in costs:
            costs[row]["monthly"][midx] = float(amt)
    return costs


def aggregate_totals(costs: dict) -> list:
    """Sum all leaf cost rows into a single 13-month series."""
    totals = [0.0] * 13
    for row, info in costs.items():
        # Only sum leaf / detail rows that have non-None budget or are in EDITABLE
        if row in EDITABLE_ROWS or (st.session_state.data["budgets"].get(row) or 0) > 0:
            # Skip pure category headers that already sum children
            if row in (12, 18, 24, 30, 39, 61):
                continue
            for i, v in enumerate(info["monthly"]):
                totals[i] += v
    # Also include any residual from parent-only amounts if children are empty
    # (simpler: use the Total Project Costs row if present, but recompute from leaves)
    return totals


def recompute_from_leaves(costs: dict) -> list:
    """Reliable total from editable / detail rows only."""
    totals = [0.0] * 13
    for row in EDITABLE_ROWS:
        if row in costs:
            for i, v in enumerate(costs[row]["monthly"]):
                totals[i] += float(v or 0)
    return totals


def compute_debt_and_cash(monthly_costs: list, rates: list):
    """
    Replicate the RMH capital + interest logic.
    Assumes all cash requirement is funded by RMH drawdowns (as in source model).
    Interest is capitalised (no interest repayments).
    """
    n = 13
    opening_cap = [0.0] * n
    drawdown = [0.0] * n
    closing_cap = [0.0] * n
    remaining = [0.0] * n
    interest = [0.0] * n
    opening_int = [0.0] * n
    closing_int = [0.0] * n
    total_rmh = [0.0] * n
    opening_cash = [0.0] * n
    closing_cash = [0.0] * n

    for i in range(n):
        if i == 0:
            opening_cap[i] = 0.0
            opening_int[i] = 0.0
            opening_cash[i] = 0.0
        else:
            opening_cap[i] = closing_cap[i - 1]
            opening_int[i] = closing_int[i - 1]
            opening_cash[i] = closing_cash[i - 1]

        # Drawdown = cash requirement for the month (rounded slightly like source)
        drawdown[i] = round(monthly_costs[i])
        closing_cap[i] = opening_cap[i] + drawdown[i]
        remaining[i] = FACILITY_LIMIT - closing_cap[i]

        # Interest on opening capital balance for the month (simple monthly)
        interest[i] = opening_cap[i] * rates[i] / 12
        closing_int[i] = opening_int[i] + interest[i]
        total_rmh[i] = closing_cap[i] + closing_int[i]

        # Cash: drawdown covers costs exactly (tiny residual from rounding)
        cash_in = drawdown[i]
        cash_out = monthly_costs[i]
        closing_cash[i] = opening_cash[i] + cash_in - cash_out

    return {
        "opening_cap": opening_cap,
        "drawdown": drawdown,
        "closing_cap": closing_cap,
        "remaining": remaining,
        "interest": interest,
        "opening_int": opening_int,
        "closing_int": closing_int,
        "total_rmh": total_rmh,
        "opening_cash": opening_cash,
        "closing_cash": closing_cash,
    }


def fmt_zar(x):
    if x is None or (isinstance(x, float) and np.isnan(x)):
        return "–"
    return f"R {x:,.0f}"


def fmt_zar_k(x):
    if x is None:
        return "–"
    if abs(x) >= 1_000_000:
        return f"R {x/1_000_000:.2f}m"
    if abs(x) >= 1_000:
        return f"R {x/1_000:.1f}k"
    return f"R {x:,.0f}"


# ---------------------------------------------------------------------------
# UI
# ---------------------------------------------------------------------------
init_session_state()

st.title("🏗️ HCRE Wellness Center")
st.caption("Cash Flow Dashboard & Actuals Tracker  ·  RMH Mezzanine Facility")

tab_dash, tab_entry, tab_detail, tab_export = st.tabs(
    ["📊 Dashboard", "✏️ Enter Actuals", "📋 Line Items", "💾 Export"]
)

costs = get_monthly_costs()
monthly_totals = recompute_from_leaves(costs)
rates = st.session_state.data["rates"]
debt = compute_debt_and_cash(monthly_totals, rates)

actual_idx = [i for i, s in enumerate(STATUS) if s == "Actual"]
forecast_idx = [i for i, s in enumerate(STATUS) if s == "Forecast"]

spent_to_date = sum(monthly_totals[i] for i in actual_idx)
forecast_remaining = sum(monthly_totals[i] for i in forecast_idx)
total_project = spent_to_date + forecast_remaining
facility_used = debt["closing_cap"][-1]
interest_accrued = debt["closing_int"][-1]
headroom = FACILITY_LIMIT - facility_used

# ---- DASHBOARD ----
with tab_dash:
    st.subheader("Key Metrics")
    c1, c2, c3, c4, c5, c6 = st.columns(6)
    c1.metric("Spent to Date", fmt_zar_k(spent_to_date), help="Jul–Oct 2026 actuals")
    c2.metric("Cost to Complete", fmt_zar_k(forecast_remaining))
    c3.metric("Total Project Cost", fmt_zar_k(total_project))
    c4.metric("Facility Drawn", fmt_zar_k(facility_used), f"{facility_used/FACILITY_LIMIT*100:.1f}% of R10m")
    c5.metric("Interest Accrued", fmt_zar_k(interest_accrued))
    c6.metric("Facility Headroom", fmt_zar_k(headroom))

    st.divider()

    # Charts row
    left, right = st.columns(2)

    with left:
        st.markdown("**Monthly Project Costs**")
        fig = go.Figure()
        colors = ["#2563eb" if s == "Actual" else "#94a3b8" for s in STATUS]
        fig.add_trace(go.Bar(
            x=MONTH_LABELS,
            y=monthly_totals,
            marker_color=colors,
            text=[fmt_zar_k(v) if v else "" for v in monthly_totals],
            textposition="outside",
            name="Costs",
        ))
        fig.update_layout(
            height=340,
            margin=dict(l=20, r=20, t=30, b=40),
            yaxis_title="R",
            showlegend=False,
            plot_bgcolor="rgba(0,0,0,0)",
            paper_bgcolor="rgba(0,0,0,0)",
        )
        st.plotly_chart(fig, use_container_width=True)
        st.caption("Blue = Actual · Grey = Forecast")

    with right:
        st.markdown("**RMH Facility Utilisation**")
        fig2 = go.Figure()
        fig2.add_trace(go.Scatter(
            x=MONTH_LABELS, y=debt["closing_cap"],
            mode="lines+markers", name="Capital drawn",
            line=dict(color="#2563eb", width=2),
        ))
        fig2.add_trace(go.Scatter(
            x=MONTH_LABELS, y=debt["total_rmh"],
            mode="lines+markers", name="Total (incl. interest)",
            line=dict(color="#7c3aed", width=2, dash="dot"),
        ))
        fig2.add_hline(y=FACILITY_LIMIT, line_dash="dash", line_color="#ef4444",
                       annotation_text="R10m facility limit")
        fig2.update_layout(
            height=340,
            margin=dict(l=20, r=20, t=30, b=40),
            yaxis_title="R",
            legend=dict(orientation="h", yanchor="bottom", y=1.02),
            plot_bgcolor="rgba(0,0,0,0)",
            paper_bgcolor="rgba(0,0,0,0)",
        )
        st.plotly_chart(fig2, use_container_width=True)

    # Second chart row
    left2, right2 = st.columns(2)

    with left2:
        st.markdown("**Cost by Category (to date + remaining)**")
        cat_data = []
        for cat, rows in CATEGORY_PARENTS.items():
            tot = 0.0
            for r in rows:
                if r in costs:
                    tot += sum(costs[r]["monthly"])
            if tot > 0:
                cat_data.append({"Category": cat, "Amount": tot})
        if cat_data:
            cdf = pd.DataFrame(cat_data)
            fig3 = px.pie(cdf, values="Amount", names="Category", hole=0.45)
            fig3.update_layout(height=320, margin=dict(l=20, r=20, t=20, b=20),
                               showlegend=True)
            st.plotly_chart(fig3, use_container_width=True)
        else:
            st.info("No costs recorded yet.")

    with right2:
        st.markdown("**Monthly Interest Charged**")
        fig4 = go.Figure()
        fig4.add_trace(go.Bar(
            x=MONTH_LABELS,
            y=debt["interest"],
            marker_color="#f59e0b",
            name="Interest",
        ))
        fig4.update_layout(
            height=320,
            margin=dict(l=20, r=20, t=30, b=40),
            yaxis_title="R",
            showlegend=False,
            plot_bgcolor="rgba(0,0,0,0)",
            paper_bgcolor="rgba(0,0,0,0)",
        )
        st.plotly_chart(fig4, use_container_width=True)

    st.divider()
    st.markdown("**Cash & Debt Snapshot (end of period)**")
    snap = pd.DataFrame({
        "Month": MONTH_LABELS,
        "Status": STATUS,
        "Project Costs": [round(v) for v in monthly_totals],
        "Drawdown": debt["drawdown"],
        "Closing Capital": [round(v) for v in debt["closing_cap"]],
        "Interest (month)": [round(v, 2) for v in debt["interest"]],
        "Closing Interest": [round(v, 2) for v in debt["closing_int"]],
        "Total RMH Balance": [round(v, 2) for v in debt["total_rmh"]],
        "Facility Remaining": [round(v) for v in debt["remaining"]],
        "Closing Cash": [round(v, 2) for v in debt["closing_cash"]],
    })
    st.dataframe(
        snap.style.format({
            "Project Costs": "R {:,.0f}",
            "Drawdown": "R {:,.0f}",
            "Closing Capital": "R {:,.0f}",
            "Interest (month)": "R {:,.2f}",
            "Closing Interest": "R {:,.2f}",
            "Total RMH Balance": "R {:,.2f}",
            "Facility Remaining": "R {:,.0f}",
            "Closing Cash": "R {:,.2f}",
        }),
        use_container_width=True,
        hide_index=True,
        height=420,
    )

# ---- ENTER ACTUALS ----
with tab_entry:
    st.subheader("Log Actual Cost / Invoice")
    st.markdown(
        "Select a line item and month, enter the amount. "
        "This **overrides** the value currently in the model for that cell and "
        "immediately refreshes the dashboard."
    )

    col_a, col_b = st.columns(2)
    with col_a:
        options = {f"{EDITABLE_ROWS[r]}  (row {r})": r for r in EDITABLE_ROWS}
        selected_label = st.selectbox("Line item", list(options.keys()))
        selected_row = options[selected_label]
    with col_b:
        month_label = st.selectbox("Month", MONTH_LABELS)
        month_idx = MONTH_LABELS.index(month_label)

    current_val = costs[selected_row]["monthly"][month_idx] if selected_row in costs else 0.0
    budget = st.session_state.data["budgets"].get(selected_row)

    st.info(
        f"Current value in model: **{fmt_zar(current_val)}**"
        + (f"  ·  Budget for this line: **{fmt_zar(budget)}**" if budget else "")
    )

    amount = st.number_input(
        "Amount (ZAR)",
        min_value=0.0,
        value=float(current_val),
        step=100.0,
        format="%.2f",
    )
    note = st.text_input("Note / invoice reference (optional)", "")

    btn_col1, btn_col2, _ = st.columns([1, 1, 3])
    with btn_col1:
        if st.button("💾 Save / Override", type="primary", use_container_width=True):
            st.session_state.overrides[(selected_row, month_idx)] = amount
            st.session_state.entries.append({
                "row": selected_row,
                "line": EDITABLE_ROWS[selected_row],
                "month": month_label,
                "month_idx": month_idx,
                "amount": amount,
                "previous": current_val,
                "note": note,
                "ts": datetime.now().strftime("%Y-%m-%d %H:%M"),
            })
            st.success(f"Updated {EDITABLE_ROWS[selected_row]} for {month_label} → {fmt_zar(amount)}")
            st.rerun()
    with btn_col2:
        if st.button("↺ Reset this cell", use_container_width=True):
            key = (selected_row, month_idx)
            if key in st.session_state.overrides:
                del st.session_state.overrides[key]
                st.success("Cell reset to original value.")
                st.rerun()
            else:
                st.warning("No override exists for this cell.")

    st.divider()
    st.markdown("**Recent entries this session**")
    if st.session_state.entries:
        edf = pd.DataFrame(st.session_state.entries[::-1])
        st.dataframe(
            edf[["ts", "line", "month", "previous", "amount", "note"]].rename(columns={
                "ts": "When", "line": "Line item", "month": "Month",
                "previous": "Previous", "amount": "New amount", "note": "Note",
            }),
            use_container_width=True,
            hide_index=True,
        )
    else:
        st.caption("No overrides recorded yet in this session.")

# ---- LINE ITEMS DETAIL ----
with tab_detail:
    st.subheader("Full Cost Line Items")
    show_zero = st.checkbox("Show zero-value rows", value=False)

    rows = []
    for row, info in sorted(costs.items()):
        name = info["name"]
        monthly = info["monthly"]
        total = sum(monthly)
        budget = st.session_state.data["budgets"].get(row)
        if not show_zero and total == 0 and (budget is None or budget == 0):
            continue
        is_override = any((row, i) in st.session_state.overrides for i in range(13))
        rows.append({
            "Row": row,
            "Category / Line": name,
            "Budget": budget,
            "Total": total,
            "Has override": "●" if is_override else "",
            **{MONTH_LABELS[i]: monthly[i] for i in range(13)},
        })

    df = pd.DataFrame(rows)
    if not df.empty:
        st.dataframe(
            df.style.format({
                "Budget": lambda x: fmt_zar(x) if pd.notna(x) else "–",
                "Total": "R {:,.0f}",
                **{m: "R {:,.0f}" for m in MONTH_LABELS},
            }, na_rep="–"),
            use_container_width=True,
            hide_index=True,
            height=520,
        )
    else:
        st.info("No line items to display.")

# ---- EXPORT ----
with tab_export:
    st.subheader("Export updated data")
    st.markdown(
        "Download a CSV of the current monthly costs (including any overrides you made), "
        "or a summary snapshot."
    )

    # Build export DF
    export_rows = []
    for row, info in sorted(costs.items()):
        if row not in EDITABLE_ROWS and row not in (12, 18, 24, 30, 39, 61):
            continue
        export_rows.append({
            "Row": row,
            "Line item": info["name"],
            "Budget": st.session_state.data["budgets"].get(row),
            **{MONTH_LABELS[i]: info["monthly"][i] for i in range(13)},
            "Total": sum(info["monthly"]),
        })
    export_df = pd.DataFrame(export_rows)

    csv_buf = io.StringIO()
    export_df.to_csv(csv_buf, index=False)
    st.download_button(
        "⬇️ Download costs CSV",
        data=csv_buf.getvalue(),
        file_name=f"HCRE_costs_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
        mime="text/csv",
    )

    # Snapshot
    snap_export = pd.DataFrame({
        "Month": MONTH_LABELS,
        "Status": STATUS,
        "Project Costs": monthly_totals,
        "RMH Drawdown": debt["drawdown"],
        "Closing Capital": debt["closing_cap"],
        "Interest charged": debt["interest"],
        "Closing Interest Bal": debt["closing_int"],
        "Total RMH Balance": debt["total_rmh"],
        "Facility Remaining": debt["remaining"],
    })
    snap_csv = io.StringIO()
    snap_export.to_csv(snap_csv, index=False)
    st.download_button(
        "⬇️ Download cash/debt snapshot CSV",
        data=snap_csv.getvalue(),
        file_name=f"HCRE_snapshot_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
        mime="text/csv",
    )

    if st.session_state.entries:
        st.markdown("**Session audit log**")
        log_df = pd.DataFrame(st.session_state.entries)
        log_csv = io.StringIO()
        log_df.to_csv(log_csv, index=False)
        st.download_button(
            "⬇️ Download entry log",
            data=log_csv.getvalue(),
            file_name=f"HCRE_entries_{datetime.now().strftime('%Y%m%d_%H%M')}.csv",
            mime="text/csv",
        )

# Sidebar
with st.sidebar:
    st.header("About")
    st.markdown(
        """
        **HCRE Wellness Center**  
        Project budget & cash-flow tracker linked to the RMH Mezzanine facility (R10 m).

        - **Dashboard** – KPIs, charts, monthly snapshot  
        - **Enter Actuals** – override any cost cell; dashboard updates live  
        - **Line Items** – full detail view  
        - **Export** – CSV downloads of current state
        """
    )
    st.divider()
    st.caption(f"Source file: `{SOURCE_XLSX.name}`")
    st.caption(f"Facility limit: R {FACILITY_LIMIT:,.0f}")
    st.caption(f"Overrides this session: {len(st.session_state.overrides)}")
    if st.button("Reset all overrides"):
        st.session_state.overrides = {}
        st.session_state.entries = []
        st.rerun()
